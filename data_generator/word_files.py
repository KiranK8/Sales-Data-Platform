"""STEP 3A: Word contracts (.docx): German free text with a small table of lanes and base rates."""
from datetime import date, timedelta

from docx import Document

from .config import INCOMING, ISSUER, make_rng
from .docs_common import FIXED_DATETIME, long_date, normalise_zip, prices_by_period, same_terms, toll_per_km
from .text import de_date, de_decimal, de_number
from .truth import DEMO_CUSTOMER

CONTRACT_COUNT = 20
TITLES = ["Preisvereinbarung", "Rahmenvertrag Transportleistungen", "Konditionenvereinbarung"]
TABLE_HEADERS = [("Strecke", "Entfernung (km)", "Grundpreis (EUR)"), ("Relation", "Entfernung (km)", "Preis in EUR")]
LANE_SEPARATORS = [" – ", " > ", " nach "]


def _percent(value, style):
    """0.18 -> '18 Prozent' / '18 %' / '18,0%' (the same value, written three ways)."""
    n = round(value * 100)
    return [f"{n} Prozent", f"{n} %", f"{n},0%"][style]


def _validity(rng, start, end, force=False, open_ended=False):
    if open_ended:
        return f"Die Preise gelten ab dem {long_date(start)} bis auf Weiteres."
    if force:
        return f"Diese Vereinbarung gilt vom {long_date(start)} bis zum {long_date(end)}."
    return rng.choice([
        f"Diese Vereinbarung gilt vom {long_date(start)} bis zum {long_date(end)}.",
        f"Gültig vom {de_date(start)} bis {de_date(end)}.",
        f"Die Preise gelten ab dem {long_date(start)} bis einschließlich {long_date(end)}.",
    ])


def _sentences(rng, charges, rate_per_km, fuel_stated, force):
    """One German sentence per surcharge. Returns [(charge_type, text)]; phrasing varies from file to file."""
    sentences = []
    if "FUEL_SURCHARGE" in charges:
        name = "Dieselzuschlag" if force else rng.choice(["Dieselzuschlag", "DSZ", "Treibstoffzuschlag"])
        pct = _percent(fuel_stated, 0 if force else rng.randrange(3))
        template = 0 if force else rng.randrange(3)
        sentences.append(("FUEL_SURCHARGE", [f"Der {name} beträgt {pct}.", f"{name}: {pct} auf den Grundpreis.",
                                             f"Als {name} werden {pct} des Grundpreises berechnet."][template]))
    if "TOLL" in charges:
        name, rate = rng.choice(["Mautzuschlag", "Maut"]), de_decimal(rate_per_km)
        sentences.append(("TOLL", rng.choice([f"Der {name} beträgt {rate} EUR pro km.", f"{name}: {rate} EUR/km.",
                                              f"Berechnet werden {rate} EUR je Kilometer ({name})."])))
    if "PEAK_SEASON" in charges:
        name, pct = rng.choice(["Hochsaisonzuschlag", "Hochsaison-Zuschlag"]), _percent(charges["PEAK_SEASON"], rng.randrange(3))
        sentences.append(("PEAK_SEASON", rng.choice([f"Der {name} beträgt {pct}.", f"{name}: {pct}."])))
    if "CURRENCY_ADJUSTMENT" in charges:
        pct = _percent(charges["CURRENCY_ADJUSTMENT"], rng.randrange(3))
        sentences.append(("CURRENCY_ADJUSTMENT", rng.choice([f"Der Währungszuschlag beträgt {pct}.", f"Währungszuschlag: {pct}."])))
    if "HANDLING" in charges:
        name, amount = rng.choice(["Umschlag", "Handling"]), de_number(charges["HANDLING"])
        sentences.append(("HANDLING", rng.choice([f"Für {name} wird eine Pauschale von {amount} EUR je Sendung berechnet.",
                                                  f"{name}: {amount} EUR pro Sendung."])))
    return sentences


def _write_contract(path, truth, customer_id, group, rng, *, demo=False, conflict=False, open_ended=False):
    """Write one contract. Returns the paragraph numbers of the interesting sentences."""
    period = group[0]
    customer = truth.customer(customer_id)["clean_name"] if demo else rng.choice(truth.aliases_of(customer_id))
    doc = Document()
    props = doc.core_properties            # fixed values: the file must be identical on every run
    props.author = props.last_modified_by = ISSUER
    props.created = props.modified = FIXED_DATETIME
    props.title = "Preisvereinbarung"

    def paragraph(text):
        doc.add_paragraph(text)
        return len(doc.paragraphs)         # 1-based paragraph number (table cells are not counted)

    doc.add_heading(f"{rng.choice(TITLES)} Nr. {customer_id}-2026", level=1)
    paragraph(f"zwischen der {ISSUER} (nachfolgend „Auftragnehmer“) und {customer} "
              "(nachfolgend „Auftraggeber“) wird Folgendes vereinbart:")
    doc.add_heading("§ 1 Geltungsdauer", level=2)
    validity_no = paragraph(_validity(rng, period.start, period.end, force=demo, open_ended=open_ended))
    doc.add_heading("§ 2 Grundpreise", level=2)
    paragraph("Für die folgenden Strecken gelten diese Grundpreise je Sendung:")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, rng.choice(TABLE_HEADERS)):
        cell.text = text
    separator = " – " if demo else rng.choice(LANE_SEPARATORS)
    for p in group:
        lane = truth.lane(p.lane_id)
        cells = table.add_row().cells
        cells[0].text = f"{lane['origin_city']}{separator}{lane['dest_city']}"
        cells[1].text = str(lane["distance_km"])
        cells[2].text = de_number(p.charges["BASE_RATE"])

    doc.add_heading("§ 3 Zuschläge", level=2)
    fuel = period.charges.get("FUEL_SURCHARGE", 0)
    fuel_stated = round(fuel + (0.02 if fuel < 0.12 else -0.02), 2) if conflict else fuel   # planted conflict
    rate = toll_per_km(truth, period) if "TOLL" in period.charges else 0
    fuel_no = None
    for charge, text in _sentences(rng, period.charges, rate, fuel_stated, demo):
        number = paragraph(text)
        if charge == "FUEL_SURCHARGE":
            fuel_no = number
    doc.add_heading("§ 4 Schlussbestimmungen", level=2)
    paragraph("Änderungen und Ergänzungen bedürfen der Schriftform.")
    paragraph(f"Hamburg, den {de_date(period.start - timedelta(days=14))}")   # signed two weeks before it starts
    doc.save(str(path))
    normalise_zip(path)
    return {"validity": validity_no, "fuel": fuel_no, "fuel_stated": fuel_stated}


def build_word_files(truth, errors, doc_prices):
    rng = make_rng("word")
    by_period = prices_by_period(truth)
    others = [c["customer_id"] for c in truth.customers if c["customer_id"] != DEMO_CUSTOMER]
    ids = sorted([DEMO_CUSTOMER] + rng.sample(others, CONTRACT_COUNT - 1))

    # which price period each contract describes
    refs = {}
    for cid in ids:
        periods = [p for p in truth.periods if p.customer_id == cid]
        # 4711's demo example is the Q4 period (valid from 1 October), not its earlier periods
        period = next(p for p in periods if p.demo and p.start == date(2026, 10, 1)) if cid == DEMO_CUSTOMER else rng.choice(periods)
        refs[cid] = (period, [period] if cid == DEMO_CUSTOMER else same_terms(truth, period))
    conflict_id = rng.choice([c for c in ids if c != DEMO_CUSTOMER and "FUEL_SURCHARGE" in refs[c][0].charges])
    open_id = rng.choice([c for c in ids if c not in (DEMO_CUSTOMER, conflict_id)])

    summary = []
    for cid in ids:
        period, group = refs[cid]
        name = f"contract_{cid}.docx"
        info = _write_contract(INCOMING / name, truth, cid, group, rng, demo=cid == DEMO_CUSTOMER,
                               conflict=cid == conflict_id, open_ended=cid == open_id)
        included = [price for p in group for price in by_period[p.pid]]
        doc_prices.extend((name, price.price_id) for price in included)
        summary.append((name, len(included)))
        fuel_price = next(price for price in included if price.charge_type == "FUEL_SURCHARGE") if cid == conflict_id else None
        if cid == conflict_id:   # PLANTED: the contract text contradicts the agreed price
            errors.add("contract_conflict", name, f"paragraph {info['fuel']}",
                       f"contract says fuel surcharge {info['fuel_stated'] * 100:.0f} %, prices.csv has "
                       f"{fuel_price.amount * 100:.0f} % ({fuel_price.price_id}, customer {cid}, lane {fuel_price.lane_id})")
        if cid == open_id:       # PLANTED: no end date in the contract
            errors.add("open_ended_contract", name, f"paragraph {info['validity']}",
                       f"valid_to missing ('bis auf Weiteres'); prices.csv has valid_to {period.end}")
    return summary
