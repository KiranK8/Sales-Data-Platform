"""Show the demo example (customer 4711, Hamburg -> Rotterdam) as it appears in each source, raw.
    python -m data_generator.show_demo
"""
import csv
import sqlite3
import sys

import openpyxl
from docx import Document

from .config import INCOMING, LEGACY, TRUTH
from .legacy_db import DB_PATH


def main():
    sys.stdout.reconfigure(encoding="utf-8")

    print("=== WORD: contract_4711.docx ===")
    doc = Document(INCOMING / "contract_4711.docx")
    for p in doc.paragraphs:
        if p.text.strip():
            print(f"  [{p.style.name}] {p.text}")
    for row in doc.tables[0].rows:
        print("  [table]", " | ".join(cell.text for cell in row.cells))

    print("\n=== EXCEL: prices_Q4_2026.xlsx (rows of customer 4711 on Hamburg -> Rotterdam) ===")
    with open(TRUTH / "customer_aliases.csv", encoding="utf-8") as f:
        names = {r["raw_name"] for r in csv.DictReader(f) if r["customer_id"] == "4711" and r["used_in"] == "files"}
    for number, row in enumerate(openpyxl.load_workbook(INCOMING / "prices_Q4_2026.xlsx").active.iter_rows(values_only=True), 1):
        if row[0] in names and row[1] and "Hamburg" in str(row[1]) and "Rotterdam" in str(row[1]):
            print(f"  row {number}: {row}")

    print("\n=== LEGACY FILE (.dat) ===")
    print("  " + "".join(str(i // 10) if i % 10 == 0 else " " for i in range(1, 87)))
    for path in sorted(INCOMING.glob("export_*.dat")):
        for line in path.read_bytes().decode("ascii").split("\r\n"):
            if line.startswith("PR4711") and "HAM-ROT" in line:
                print(f"  {line}   <- {path.name}")

    print("\n=== LEGACY DATABASE: customers 4711 and its duplicate record 9003 ===")
    con = sqlite3.connect(DB_PATH)
    for row in con.execute("SELECT * FROM customers WHERE id IN (4711, 9003)"):
        print(" ", row)
    con.close()


if __name__ == "__main__":
    main()
